"""
Excel parsing utilities.

Client sheet: fully user-configurable (header row, data start row, tag column).
INFORM sheet: fixed 3-row header format —
    Row 1  : Column headers (may use merged cells spanning Original + New column pairs)
    Row 2  : Unrequired metadata (skipped)
    Row 3  : "Original" / "New" marker for each physical column
    Row 4+ : Data rows
Only columns whose Row-3 marker is "Original" are extracted.
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_sheet_names(file_bytes: bytes) -> list[str]:
    """Return all sheet names from an Excel file."""
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def _build_merged_cell_map(ws) -> dict[tuple[int, int], object]:
    """
    Return a dict mapping (row, col) → value for every cell that belongs to
    a merged range, using the top-left cell's value for all cells in the range.
    """
    merged_map: dict[tuple[int, int], object] = {}
    for merge_range in ws.merged_cells.ranges:
        top_left_val = ws.cell(merge_range.min_row, merge_range.min_col).value
        for row in range(merge_range.min_row, merge_range.max_row + 1):
            for col in range(merge_range.min_col, merge_range.max_col + 1):
                merged_map[(row, col)] = top_left_val
    return merged_map


def _unique_header(name: str, seen: dict[str, int]) -> str:
    """Return a deduplicated header name (appends .1, .2 … on collision)."""
    if name not in seen:
        seen[name] = 0
        return name
    seen[name] += 1
    return f"{name}.{seen[name]}"


# ---------------------------------------------------------------------------
# INFORM sheet parser
# ---------------------------------------------------------------------------

def parse_inform_sheet(
    file_bytes: bytes,
    sheet_name: str,
) -> tuple[pd.DataFrame, list[str]]:
    """
    Parse an INFORM sheet using its fixed 3-row header format.

    Returns
    -------
    df : pd.DataFrame
        Contains only the "Original" columns, with data from row 4 onwards.
    original_headers : list[str]
        Ordered list of Original column header names (as they appear in the DataFrame).
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb[sheet_name]

    merged_map = _build_merged_cell_map(ws)

    def cell_val(row: int, col: int):
        return merged_map.get((row, col), ws.cell(row, col).value)

    max_col = ws.max_column

    # Row 1 → raw header names; Row 3 → Original / New markers
    headers_row1 = [cell_val(1, c) for c in range(1, max_col + 1)]
    markers_row3 = [cell_val(3, c) for c in range(1, max_col + 1)]

    # Identify Original columns (0-indexed positions for data slicing)
    original_positions: list[int] = []
    original_headers: list[str] = []
    seen: dict[str, int] = {}

    for i, (hdr, marker) in enumerate(zip(headers_row1, markers_row3)):
        if marker and str(marker).strip().lower() == "original" and hdr is not None:
            clean = _unique_header(str(hdr).strip(), seen)
            original_positions.append(i)
            original_headers.append(clean)

    # Read data from row 4 onwards — only Original column positions
    data: list[list] = []
    for row_tuple in ws.iter_rows(min_row=4, values_only=True):
        row_data = [row_tuple[i] for i in original_positions]
        if any(v is not None for v in row_data):
            data.append(row_data)

    wb.close()

    df = pd.DataFrame(data, columns=original_headers)
    return df, original_headers


# ---------------------------------------------------------------------------
# Client sheet parser
# ---------------------------------------------------------------------------

def parse_client_sheet(
    file_bytes: bytes,
    sheet_name: str,
    header_row: int,
    data_start_row: int,
) -> tuple[pd.DataFrame, list[str]]:
    """
    Parse a client sheet using user-supplied row configuration.

    Parameters
    ----------
    header_row : int
        1-indexed row number that contains column headers.
    data_start_row : int
        1-indexed row number where data begins (must be > header_row).

    Returns
    -------
    df : pd.DataFrame
        Full data frame of the client sheet.
    headers : list[str]
        Ordered list of column header names.
    """
    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb[sheet_name]

    max_col = ws.max_column

    # Read headers from the specified row
    raw_headers = [ws.cell(header_row, c).value for c in range(1, max_col + 1)]

    seen: dict[str, int] = {}
    cleaned_headers: list[str] = []
    for h in raw_headers:
        name = str(h).strip() if h is not None else f"_col_{len(cleaned_headers) + 1}"
        cleaned_headers.append(_unique_header(name, seen))

    # Read data rows
    data: list[list] = []
    for row_tuple in ws.iter_rows(min_row=data_start_row, values_only=True):
        if any(v is not None for v in row_tuple):
            data.append(list(row_tuple))

    wb.close()

    df = pd.DataFrame(data, columns=cleaned_headers)
    return df, cleaned_headers


# ---------------------------------------------------------------------------
# Preview helper
# ---------------------------------------------------------------------------

def preview_rows(
    file_bytes: bytes,
    sheet_name: str,
    start_row: int,
    n_rows: int = 8,
) -> pd.DataFrame:
    """
    Return a raw preview of `n_rows` rows starting at `start_row` (1-indexed).
    Used to help users identify header / data rows visually.
    """
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb[sheet_name]

    rows = []
    end_row = start_row + n_rows - 1
    for row_tuple in ws.iter_rows(
        min_row=start_row, max_row=end_row, values_only=True
    ):
        rows.append(list(row_tuple))

    wb.close()

    if not rows:
        return pd.DataFrame()

    max_cols = max(len(r) for r in rows)
    padded = [r + [None] * (max_cols - len(r)) for r in rows]
    col_labels = [f"Col {i + 1}" for i in range(max_cols)]
    df = pd.DataFrame(padded, columns=col_labels)
    df.index = range(start_row, start_row + len(df))
    return df
