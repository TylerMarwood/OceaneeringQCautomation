"""
Excel report generation.

Two reports are available:
  1. Tag Issues Report   — produced whenever tag issues exist (duplicates /
                           one-sided tags).  Can also be downloaded when tags
                           are clean (will state no issues found).
  2. Full Comparison Report — multi-sheet workbook:
       • Summary       — key statistics
       • Tag Issues    — all tag-level problems (if any)
       • Mismatches    — only the differing cells
       • All Results   — every compared cell (green = match, red = mismatch)
       • Header Info   — matched / unmatched headers
"""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import openpyxl
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from .comparison import ComparisonResult


# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

_GREEN_FILL   = PatternFill("solid", fgColor="C6EFCE")
_RED_FILL     = PatternFill("solid", fgColor="FFC7CE")
_AMBER_FILL   = PatternFill("solid", fgColor="FFEB9C")
_HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
_TITLE_FILL   = PatternFill("solid", fgColor="2E75B6")
_SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")

_HEADER_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_TITLE_FONT   = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
_BOLD_FONT    = Font(name="Calibri", bold=True, size=11)
_BODY_FONT    = Font(name="Calibri", size=11)

_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

_THIN  = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _col_width(ws, col_letter: str, width: float) -> None:
    ws.column_dimensions[col_letter].width = width


def _auto_width(ws, min_w: int = 10, max_w: int = 50) -> None:
    """Set column widths based on maximum cell content length."""
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        best = min_w
        for cell in col_cells:
            if cell.value:
                best = max(best, min(len(str(cell.value)) + 4, max_w))
        ws.column_dimensions[letter].width = best


def _title_row(ws, text: str, ncols: int, row: int = 1) -> None:
    """Write a full-width title row."""
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=ncols
    )
    cell = ws.cell(row, 1, text)
    cell.font  = _TITLE_FONT
    cell.fill  = _TITLE_FILL
    cell.alignment = _CENTER
    ws.row_dimensions[row].height = 28


def _header_row(ws, headers: list[str], row: int = 2) -> None:
    """Write a styled header row."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border    = _BORDER


def _write_df(
    ws,
    df: pd.DataFrame,
    start_row: int = 3,
    row_fill_fn=None,
) -> int:
    """
    Write *df* to worksheet starting at *start_row*.
    Optional *row_fill_fn(row_dict) -> PatternFill | None* colours each row.
    Returns the last data row written.
    """
    for r_offset, (_, row_data) in enumerate(df.iterrows()):
        r = start_row + r_offset
        fill = row_fill_fn(row_data) if row_fill_fn else None
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(r, c, "" if pd.isna(val) else val)
            cell.font      = _BODY_FONT
            cell.alignment = _LEFT
            cell.border    = _BORDER
            if fill:
                cell.fill = fill
    return start_row + len(df) - 1


# ---------------------------------------------------------------------------
# Report 1: Tag Issues
# ---------------------------------------------------------------------------

def generate_tag_issues_report(issues_df: pd.DataFrame) -> bytes:
    """
    Produce a standalone Excel report for tag-level issues.
    If *issues_df* is empty, the report states no issues were found.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tag Issues"

    cols = ["Tag Number", "Sheet", "Reason"]
    _title_row(ws, "Tag Comparison — Issues Report", len(cols), row=1)
    _header_row(ws, cols, row=2)

    if issues_df.empty:
        ws.merge_cells("A3:C3")
        cell = ws.cell(3, 1, "No tag issues found — all tags are unique and present in both sheets.")
        cell.font      = Font(name="Calibri", italic=True, size=11)
        cell.fill      = _GREEN_FILL
        cell.alignment = _CENTER
        cell.border    = _BORDER
    else:
        display_df = issues_df.rename(
            columns={"Tag": "Tag Number", "Sheet": "Sheet", "Reason": "Reason"}
        )

        def row_fill(row):
            reason = str(row.get("Reason", ""))
            if "Duplicate" in reason:
                return _RED_FILL
            return _AMBER_FILL

        _write_df(ws, display_df[cols], start_row=3, row_fill_fn=row_fill)

    _auto_width(ws)
    ws.freeze_panes = "A3"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Report 2: Full Comparison
# ---------------------------------------------------------------------------

def generate_comparison_report(result: ComparisonResult) -> bytes:
    """
    Multi-sheet workbook containing:
      Summary | Tag Issues | Header Info | Mismatches | All Results
    """
    wb = openpyxl.Workbook()

    # -----------------------------------------------------------------------
    # Sheet 1: Summary
    # -----------------------------------------------------------------------
    ws_s = wb.active
    ws_s.title = "Summary"

    tv = result.tag_validation
    hm = result.header_match
    df_diffs = result.diffs_df()

    total      = len(df_diffs)
    matches    = int(result.match_count)
    mismatches = int(result.mismatch_count)
    rate       = f"{matches / total * 100:.1f} %" if total else "N/A"

    summary_rows = [
        ("QC Comparison Summary", ""),
        ("", ""),
        ("DATA COMPARISON", ""),
        ("Total cell comparisons",           total),
        ("Matching cells",                   matches),
        ("Mismatching cells",                mismatches),
        ("Match rate",                       rate),
        ("", ""),
        ("TAG VALIDATION", ""),
        ("Comparable tags (both sheets)",    len(tv.comparable_tags)),
        ("Client-only tags",                 len(tv.client_only_tags)),
        ("INFORM-only tags",                 len(tv.inform_only_tags)),
        ("Tag issues total",                 len(tv.issues)),
        ("", ""),
        ("HEADER MATCHING", ""),
        ("Matched headers",                  len(hm.matched)),
        ("Client-only headers",              len(hm.client_only)),
        ("INFORM-only headers",              len(hm.inform_only)),
    ]

    for r, (label, value) in enumerate(summary_rows, 1):
        label_cell = ws_s.cell(r, 1, label)
        value_cell = ws_s.cell(r, 2, value)

        if label in ("QC Comparison Summary",):
            label_cell.font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
            label_cell.fill = _TITLE_FILL
            ws_s.merge_cells(f"A{r}:B{r}")
            label_cell.alignment = _CENTER
        elif label in ("DATA COMPARISON", "TAG VALIDATION", "HEADER MATCHING"):
            for cell in (label_cell, value_cell):
                cell.font  = _BOLD_FONT
                cell.fill  = _SECTION_FILL
                cell.border = _BORDER
        elif label:
            label_cell.font  = _BODY_FONT
            label_cell.border = _BORDER
            value_cell.font  = _BODY_FONT
            value_cell.border = _BORDER

            # Colour bad metrics red
            if isinstance(value, int) and label in (
                "Mismatching cells", "Client-only tags",
                "INFORM-only tags", "Tag issues total",
                "Client-only headers", "INFORM-only headers",
            ) and value > 0:
                value_cell.fill = _RED_FILL
            elif isinstance(value, int) and label in (
                "Matching cells", "Comparable tags (both sheets)", "Matched headers"
            ) and value > 0:
                value_cell.fill = _GREEN_FILL

    ws_s.column_dimensions["A"].width = 38
    ws_s.column_dimensions["B"].width = 20

    # -----------------------------------------------------------------------
    # Sheet 2: Tag Issues
    # -----------------------------------------------------------------------
    ws_t = wb.create_sheet("Tag Issues")
    issues_df = tv.issues_df()
    cols_t = ["Tag", "Sheet", "Reason"]
    _title_row(ws_t, "Tag Issues", len(cols_t), row=1)
    _header_row(ws_t, cols_t, row=2)

    if issues_df.empty:
        ws_t.merge_cells("A3:C3")
        cell = ws_t.cell(3, 1, "No tag issues — all tags are unique and present in both sheets.")
        cell.font      = Font(name="Calibri", italic=True, size=11)
        cell.fill      = _GREEN_FILL
        cell.alignment = _CENTER
        cell.border    = _BORDER
    else:
        def tag_row_fill(row):
            return _RED_FILL if "Duplicate" in str(row.get("Reason", "")) else _AMBER_FILL
        _write_df(ws_t, issues_df[cols_t], start_row=3, row_fill_fn=tag_row_fill)

    _auto_width(ws_t)
    ws_t.freeze_panes = "A3"

    # -----------------------------------------------------------------------
    # Sheet 3: Header Info
    # -----------------------------------------------------------------------
    ws_h = wb.create_sheet("Header Info")
    max_len = max(len(hm.matched), len(hm.client_only), len(hm.inform_only), 1)
    _title_row(ws_h, "Header Matching Results", 3, row=1)
    _header_row(ws_h, ["Matched Headers", "Client-Only Headers", "INFORM-Only Headers"], row=2)

    for r in range(max_len):
        m  = hm.matched[r]     if r < len(hm.matched)     else ""
        co = hm.client_only[r] if r < len(hm.client_only) else ""
        io = hm.inform_only[r] if r < len(hm.inform_only) else ""

        c1 = ws_h.cell(r + 3, 1, m)
        c2 = ws_h.cell(r + 3, 2, co)
        c3 = ws_h.cell(r + 3, 3, io)

        for cell, fill in ((c1, _GREEN_FILL if m else None),
                           (c2, _AMBER_FILL if co else None),
                           (c3, _AMBER_FILL if io else None)):
            cell.font      = _BODY_FONT
            cell.alignment = _LEFT
            cell.border    = _BORDER
            if fill:
                cell.fill = fill

    _auto_width(ws_h)
    ws_h.freeze_panes = "A3"

    # -----------------------------------------------------------------------
    # Sheet 4: Mismatches only
    # -----------------------------------------------------------------------
    ws_m = wb.create_sheet("Mismatches")
    mismatch_df = df_diffs[~df_diffs["Match"]] if not df_diffs.empty else pd.DataFrame()
    cols_r = ["Tag", "Header", "Client Value", "INFORM Value",
              "Client (Normalised)", "INFORM (Normalised)", "Match"]

    _title_row(ws_m, "Mismatching Cells", len(cols_r), row=1)
    _header_row(ws_m, cols_r, row=2)

    if mismatch_df.empty:
        ws_m.merge_cells(f"A3:{get_column_letter(len(cols_r))}3")
        cell = ws_m.cell(3, 1, "No mismatches found — all compared cells match.")
        cell.font      = Font(name="Calibri", italic=True, size=11)
        cell.fill      = _GREEN_FILL
        cell.alignment = _CENTER
        cell.border    = _BORDER
    else:
        _write_df(ws_m, mismatch_df[cols_r], start_row=3,
                  row_fill_fn=lambda _: _RED_FILL)

    _auto_width(ws_m, max_w=40)
    ws_m.freeze_panes = "A3"

    # -----------------------------------------------------------------------
    # Sheet 5: All Results
    # -----------------------------------------------------------------------
    ws_a = wb.create_sheet("All Results")
    _title_row(ws_a, "Full Comparison — All Cells", len(cols_r), row=1)
    _header_row(ws_a, cols_r, row=2)

    if df_diffs.empty:
        ws_a.merge_cells(f"A3:{get_column_letter(len(cols_r))}3")
        cell = ws_a.cell(3, 1, "No comparison data available.")
        cell.font      = Font(name="Calibri", italic=True, size=11)
        cell.alignment = _CENTER
        cell.border    = _BORDER
    else:
        def all_row_fill(row):
            return _GREEN_FILL if row["Match"] else _RED_FILL
        _write_df(ws_a, df_diffs[cols_r], start_row=3, row_fill_fn=all_row_fill)

    _auto_width(ws_a, max_w=40)
    ws_a.freeze_panes = "A3"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
